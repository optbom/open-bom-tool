"""Read CSV/XLSX BOM exports, detect columns, and normalize row values."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Mapping

from ..errors import OpenBomError
from ..objects import QUANTITY_RE

HEADER_SCAN_ROWS = 25

# Canonical fields and the header spellings that map to them (compared after _norm_header).
SYNONYMS: dict[str, tuple[str, ...]] = {
    "level": ("level", "lvl", "bom level", "indent", "indent level", "indented level"),
    "part_number": (
        "part number", "part no", "part #", "part", "pn", "p/n", "part id", "item number", "item no",
        "item #", "item", "number", "component", "component number",
    ),
    "revision": ("rev", "revision", "item rev", "item revision", "part rev", "part revision"),
    "name": ("name", "item name", "part name", "short description", "title"),
    "description": ("description", "desc", "item description", "part description", "long description"),
    "quantity": ("qty", "quantity", "qty per", "quantity per", "qty/assy", "qty per assembly", "qty per assy"),
    "uom": ("uom", "unit", "units", "unit of measure", "u/m"),
    "refdes": (
        "ref des", "refdes", "reference designator", "reference designators", "ref designator",
        "ref designators", "designator", "designators",
    ),
    "manufacturer": ("manufacturer", "mfr", "mfg", "mfr name", "mfg name", "manufacturer name"),
    "customer_pn": ("customer pn", "customer p/n", "customer part number", "customer part no", "cpn"),
    "mpn": (
        "mpn", "mfr p/n", "mfr pn", "mfr part number", "mfr part #", "mfg p/n", "mfg pn", "mfg part number",
        "manufacturer part number", "manufacturer p/n",
    ),
}
FIELDS = tuple(SYNONYMS)
REQUIRED_FIELDS = ("level", "part_number", "quantity")

UOM_ALIASES = {"each": "ea", "pcs": "ea", "pc": "ea", "piece": "ea", "pieces": "ea", "ea.": "ea", "qty": "ea"}


def _norm_header(text: str) -> str:
    text = text.lower().replace("_", " ").replace("-", " ")
    text = re.sub(r"[^a-z0-9#/ ]", "", text)
    return " ".join(text.split())


_SYNONYM_INDEX = {_norm_header(s): f for f, spellings in SYNONYMS.items() for s in spellings}


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------


@dataclass
class Table:
    path: Path
    rows: list[tuple[int, list[str]]]  # (1-based row number in the file, cell strings)
    sheet: str | None = None


def cell_text(value: Any) -> str:
    """Spreadsheet cell -> string, without float artifacts (1.0 -> "1")."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def read_table(path: Path, sheet: str | None = None) -> Table:
    if not path.exists():
        raise OpenBomError(f"file {path} does not exist")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet)
    if suffix in (".csv", ".tsv", ".txt"):
        if sheet is not None:
            raise OpenBomError("--sheet only applies to .xlsx files")
        return _read_csv(path)
    raise OpenBomError(f"unsupported file type {suffix!r}: use .csv, .tsv, or .xlsx")


def _read_csv(path: Path) -> Table:
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = [(i, [c.strip() for c in r]) for i, r in enumerate(csv.reader(text.splitlines(), dialect), 1)]
    return Table(path=path, rows=rows)


def _read_xlsx(path: Path, sheet: str | None) -> Table:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise OpenBomError("reading .xlsx needs openpyxl (pip install openpyxl)") from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise OpenBomError(f"sheet {sheet!r} not found; sheets: {', '.join(wb.sheetnames)}")
        ws = wb[sheet] if sheet is not None else wb.worksheets[0]
        rows = [(i, [cell_text(v) for v in r]) for i, r in enumerate(ws.iter_rows(values_only=True), 1)]
        return Table(path=path, rows=rows, sheet=ws.title)
    finally:
        wb.close()


# --------------------------------------------------------------------------
# Column mapping
# --------------------------------------------------------------------------


@dataclass
class ColumnMap:
    header_row: int
    headers: list[str]
    columns: dict[str, int]  # field -> column index
    ignored: list[str] = field(default_factory=list)  # header cells not mapped to a field

    def describe(self) -> dict[str, str]:
        return {f: self.headers[i] for f, i in self.columns.items()}


def detect_columns(table: Table, overrides: Mapping[str, str] | None = None) -> ColumnMap:
    """Find the header row (exports often have title rows above it) and map columns to fields.

    overrides: field -> exact header text, for headers the synonyms don't cover.
    """
    overrides = dict(overrides or {})
    for f in overrides:
        if f not in FIELDS:
            raise OpenBomError(f"unknown field {f!r} in --column; fields are: {', '.join(FIELDS)}")
    override_index = {_norm_header(h): f for f, h in overrides.items()}

    best: ColumnMap | None = None
    for row_no, cells in table.rows[:HEADER_SCAN_ROWS]:
        columns: dict[str, int] = {}
        ignored = []
        for i, cell in enumerate(cells):
            key = _norm_header(cell)
            f = override_index.get(key) or _SYNONYM_INDEX.get(key)
            if f is not None and f not in columns:
                columns[f] = i
            elif cell:
                ignored.append(cell)
        if best is None or len(columns) > len(best.columns):
            best = ColumnMap(header_row=row_no, headers=cells, columns=columns, ignored=ignored)
    if best is None or "part_number" not in best.columns:
        raise OpenBomError(
            f"could not find a header row with a part number column in the first {HEADER_SCAN_ROWS} rows; "
            "use --column part_number=\"<header>\""
        )
    missing = [f for f in REQUIRED_FIELDS if f not in best.columns]
    if missing:
        raise OpenBomError(
            f"header row {best.header_row} has no {', '.join(missing)} column "
            f"(found: {', '.join(f'{f}={h!r}' for f, h in best.describe().items())}); "
            "use --column <field>=\"<header>\""
        )
    return best


# --------------------------------------------------------------------------
# Value normalization
# --------------------------------------------------------------------------


def parse_level(raw: str) -> int | None:
    """'2', '.2', '..2', '..' (dot count) -> int; None if unparseable."""
    raw = raw.strip()
    if raw.isdigit():
        return int(raw)
    m = re.fullmatch(r"(\.+)([0-9]*)", raw)
    if m:
        return int(m.group(2)) if m.group(2) else len(m.group(1))
    return None


def normalize_quantity(raw: str) -> str | None:
    """'1.0' -> '1', '1,000' -> '1000', '0.50' -> '0.5'; None if not a positive number."""
    try:
        d = Decimal(raw.replace(",", "").strip())
    except InvalidOperation:
        return None
    if not d.is_finite() or d <= 0:
        return None
    text = format(d.normalize(), "f")
    return text if QUANTITY_RE.fullmatch(text) else None


def normalize_uom(raw: str) -> str | None:
    raw = raw.strip().lower()
    return UOM_ALIASES.get(raw, raw) or None


_RANGE_RE = re.compile(r"([A-Za-z_]+)([0-9]+)-(?:\1)?([0-9]+)")
MAX_RANGE = 10_000


def parse_refdes(raw: str) -> list[str]:
    """'C1-C4, R7' or 'C1-4 R7' -> ['C1', 'C2', 'C3', 'C4', 'R7']."""
    result: list[str] = []
    for token in re.split(r"[,;\s]+", raw.strip()):
        if not token:
            continue
        m = _RANGE_RE.fullmatch(token)
        if m:
            prefix, start, end = m.group(1), int(m.group(2)), int(m.group(3))
            if start <= end and end - start < MAX_RANGE:
                result.extend(f"{prefix}{n}" for n in range(start, end + 1))
                continue
        result.append(token)
    return result


@dataclass
class Row:
    line: int
    level: int | None
    pn: str
    raw_level: str
    raw_quantity: str
    quantity: str | None
    revision: str | None = None
    name: str | None = None
    description: str | None = None
    uom: str | None = None
    refdes: tuple[str, ...] = ()
    manufacturer: str | None = None
    mpn: str | None = None
    # Approved manufacturer list, when the source gives more than one (manufacturer, mpn).
    aml: tuple[tuple[str, str], ...] = ()
    # Cross-references such as ("customer", "01-950837") or ("supplier", "9Z34A900-600-G").
    identifiers: tuple[tuple[str, str], ...] = ()
    # Line-level substitutes folded in by the profile: (priority, row).
    substitutes: list[tuple[int | None, "Row"]] = field(default_factory=list)


def split_values(raw: str) -> list[str]:
    """A cell holding several values ('19-974557\n19-958107', 'A, B') -> values, minus placeholders."""
    return [v for v in (x.strip() for x in re.split(r"[\n,;]+", raw)) if v and not is_placeholder(v)]


PLACEHOLDERS = {"NA", "N/A", "TBD", "TBA", "-", "NONE", "NULL"}


def is_placeholder(value: str) -> bool:
    return value.strip().upper() in PLACEHOLDERS


def clean(value: str) -> str:
    """Strip tabs, non-breaking spaces, and placeholder values like NA/TBD."""
    value = value.replace("\xa0", " ").strip()
    return "" if is_placeholder(value) else value


def read_rows(table: Table, columns: ColumnMap) -> list[Row]:
    """Data rows below the header, skipping blank rows. Values are normalized but not validated."""

    def get(cells: list[str], f: str) -> str:
        i = columns.columns.get(f)
        return cells[i].strip() if i is not None and i < len(cells) else ""

    rows = []
    for line, cells in table.rows:
        if line <= columns.header_row or not any(c.strip() for c in cells):
            continue
        raw_level, raw_qty = get(cells, "level"), get(cells, "quantity")
        name, description = get(cells, "name"), get(cells, "description")
        if not name and description and "name" not in columns.columns:
            name, description = description, ""  # a lone description column is the part name
        rows.append(
            Row(
                line=line,
                level=parse_level(raw_level),
                pn=get(cells, "part_number"),
                raw_level=raw_level,
                raw_quantity=raw_qty,
                quantity=normalize_quantity(raw_qty) if raw_qty else None,
                revision=get(cells, "revision") or None,
                name=name or None,
                description=description or None,
                uom=normalize_uom(get(cells, "uom")),
                refdes=tuple(parse_refdes(get(cells, "refdes"))),
                manufacturer=get(cells, "manufacturer") or None,
                mpn=get(cells, "mpn") or None,
                identifiers=tuple(("customer", v) for v in split_values(get(cells, "customer_pn"))),
            )
        )
    return rows
