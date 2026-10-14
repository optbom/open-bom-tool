"""Make a synthetic "next export" of an Agile BOM to demo versioned reconciliation.

Usage: python examples/make_agile_rev2.py IN.xlsx OUT.xlsx

Applies three ODM-side changes, as an engineering change would:
  1. a data SSD quantity in the Bay B (left) module goes 2 -> 3
  2. the HPM quantity on the tray goes 2 -> 1
  3. the PDB-to-NF4 power cable switches to a different MPN (and its rev bumps)
"""

import sys

import openpyxl

CHANGES = [
    # (part number, current qty, field, new value)
    ("MTTEMP-S927101354", "2", "BOM.Qty", "3"),
    ("MTTEMP-S987539199", "2", "BOM.Qty", "1"),
    ("MTTEMP-S961902666", None, "Manufacturers.Mfr Part Number", "MPN542-16370"),
    ("MTTEMP-S961902666", None, "Rev", "006        ECSYNZ99999999"),
]


def main(src: str, dst: str) -> None:
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    header = {str(c.value).split("/")[0].strip(): c.column for c in ws[1] if c.value}
    for number, qty, field, value in CHANGES:
        for row in ws.iter_rows(min_row=2):
            if str(row[header["Number"] - 1].value or "").strip() != number:
                continue
            if qty is not None and str(row[header["BOM.Qty"] - 1].value).strip() != qty:
                continue
            row[header[field] - 1].value = value
            print(f"row {row[0].row}: {number} {field} -> {value}")
            break
        else:
            raise SystemExit(f"no row for {number} (qty {qty})")
    wb.save(dst)


if __name__ == "__main__":
    main(*sys.argv[1:3])
